"""Journal des trades VÉCUS EN DIRECT (Strat5669 / Bot5669), reconstruit depuis le journal d'alertes TradingView,
puis comparé à l'historique recalculé (export xlsx) pour compter les fantômes.

Le graphique recalcule tout à chaque rechargement : un trade pris en direct peut disparaître ou changer. Le journal
d'alertes, lui, ne bouge plus : c'est la seule trace de ce qui a vraiment été signalé. Les statistiques « en direct »
ci-dessous sont donc celles qui comptent pour juger la méthode ; celles du backtest sont un repère.

Usage : uv run --with pandas --with openpyxl python verif_fantomes.py <dossier KAIROS-data> [--csv sortie.csv]
Cherche récursivement les TradingView_Alerts_Log*.csv et les exports Strat5669 / Stratégie 56 / Bot5669.
"""
import glob, os, re, sys
import pandas as pd

args = [a for a in sys.argv[1:] if not a.startswith("--")]
racine = args[0] if args else "."
sortie = sys.argv[sys.argv.index("--csv") + 1] if "--csv" in sys.argv else None
verrou = sys.argv[sys.argv.index("--verrou") + 1] if "--verrou" in sys.argv else None
suivi = sys.argv[sys.argv.index("--suivi") + 1] if "--suivi" in sys.argv else None


def fichiers(motif):
    return glob.glob(os.path.join(glob.escape(racine), "**", motif), recursive=True)


def nombre(txt, motif):
    m = re.search(motif, txt)
    return float(m.group(1).replace(",", "").replace(" ", "").replace(" ", "")) if m else None


# --- 1. alertes ------------------------------------------------------------------------------------------------
logs = pd.concat([pd.read_csv(f).assign(fichier=f) for f in fichiers("TradingView_Alerts_Log*.csv")])
logs = logs.drop_duplicates(["ID de l'alerte", "Heure", "Description"])
logs["t"] = pd.to_datetime(logs["Heure"]).dt.tz_convert("Europe/Paris").dt.tz_localize(None)
logs["sym"] = logs["Ticker"].str.split(",").str[0].str.split(":").str[-1]
logs["tf"] = logs["Ticker"].str.split(", ").str[1]
logs = logs.sort_values("t", kind="stable")
# heure du dernier message de chaque journal ≈ heure de son export : une alerte encore active a tourné jusque-là
logs["fin_journal"] = logs.groupby("fichier")["t"].transform("max")
# nom de la stratégie de chaque alerte (lu dans les messages « ordre … effectué » qui commencent par le nom)
STRATS = ("correctifstest", "Bot5669", "Strat5669", "Stratégie 56")
strat_id = {}
for _, a in logs.iterrows():
    for n in STRATS:
        if str(a["Description"]).startswith(n):
            strat_id[a["ID de l'alerte"]] = n

trades, ouverts = [], {}
for _, a in logs.iterrows():
    d, cle = str(a["Description"]), (a["ID de l'alerte"], a["sym"])
    if "au marché maintenant" in d and re.search("GODLIKE|56 seul", d):
        ouverts[cle] = dict(id=a["ID de l'alerte"], strat=strat_id.get(a["ID de l'alerte"], "Strat5669"), entree=a.t, sym=a.sym, tf=a.tf, sens="L" if "ACHAT" in d else "S",
                            etiq="GODLIKE" if "GODLIKE" in d else "56 seul",
                            px=nombre(d, r"Entrée ≈ ([\d.,]+)"), stop=nombre(d, r"Stop ([\d.,]+)"),
                            obj=nombre(d, r"Objectif ([\d.,]+)"), sortie=None, raison=None, px_sortie=None, pnl=None)
        trades.append(ouverts[cle])
    elif cle in ouverts and ("trade terminé" in d or "FERMER la position" in d):
        tr = ouverts.pop(cle)
        tr["sortie"] = a.t
        if "trade terminé" in d:
            tr["raison"] = "objectif" if "OBJECTIF" in d else "stop"
            tr["px_sortie"] = nombre(d, r"touché à ([\d.,]+)")
            tr["pnl"] = nombre(d, r"P&L (-?[\d.,  ]+) \$")
        else:
            tr["raison"] = "fin de fenêtre" if "fin de la fenêtre" in d else "temps"
            tr["px_sortie"] = nombre(d, r"prix ≈ ([\d.,]+)")
            tr["pnl"] = nombre(d, r"P&L ≈ (-?[\d.,  ]+) \$")

# --- 2. historique recalculé -------------------------------------------------------------------------------------
exports = {}
for f in fichiers("*.xlsx"):
    nom = os.path.basename(f)
    m = re.search(r"(MNQ1!|MES1!|BTCUSD|BTC1!|MBT1!)", nom)
    st = re.match(r"(correctifstest|Strat5669|Stratégie_56|Bot5669)", nom)
    if not m or not st:
        continue
    st = {"Stratégie_56": "Strat5669"}.get(st.group(1), st.group(1))
    tf = "5m" if (os.sep + "5min" + os.sep in f or nom.startswith("Bot")) else "15m"
    t = pd.read_excel(f, "Trades")
    t["Date et heure"] = pd.to_datetime(t["Date et heure"])
    exports.setdefault((st, m.group(1), tf), []).append((t["Date et heure"].max(), t))

for tr in trades:
    couverts = [t for mx, t in exports.get((tr["strat"] if tr["strat"] != "Stratégie 56" else "Strat5669", tr["sym"], tr["tf"]), []) if tr["entree"] <= mx]
    tr["historique"], tr["pnl_hist"] = "pas encore exporté", None
    if not couverts:
        continue
    t = max(couverts, key=len)
    e = t[t["Type"].str.startswith("Entrer") & ((t["Date et heure"] - tr["entree"]).abs() <= pd.Timedelta(minutes=15))]
    e = e[e["Signal"].str.startswith(tr["sens"])]
    if e.empty:
        tr["historique"] = "ABSENT (fantôme)"
        continue
    r = e.iloc[0]
    etiq = "GODLIKE" if "GODLIKE" in r["Signal"] else "56 seul"
    tr["historique"] = "présent" if etiq == tr["etiq"] else f"présent, devenu « {etiq} »"
    tr["pnl_hist"] = r["P&L net USD"]
    if tr["pnl"] is None and tr["sortie"] is not None:
        tr["pnl"], tr["raison"] = r["P&L net USD"], tr["raison"] + " (P&L de l'historique)"

# --- fiche de suivi (xlsx, une feuille par actif) -----------------------------------------------------------------
FEUILLES = {"MNQ": ("MNQ1!",), "MES": ("MES1!",), "BTC": ("BTC1!", "MBT1!", "BTCUSD")}
AUTO = ["Date", "Heure", "Sens", "Étiquette", "Résultat", "P&L alerte $", "Fantôme ?"]
PERSO = ["Pris ?", "Si non, pourquoi", "Mon entrée", "Ma sortie", "Mon P&L $", "P&L FTMO $", "Remarque"]
RAISONS = "absent,risque trop gros,entrée trop tardive,56 seul,oublié / pas vu,autre"


def ecrire_suivi(df, chemin):
    """Fiche à remplir par Jeremy. Colonnes grises = remplies depuis les alertes ; colonnes jaunes = les siennes,
    jamais écrasées (relues dans l'ancien fichier, clé = feuille + symbole + date + heure + sens)."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule

    anciens = {}
    if os.path.exists(chemin):
        wb0 = load_workbook(chemin)
        for ws in wb0.worksheets:
            tete = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                r = dict(zip(tete, row))
                if r.get("Date"):
                    dt, hr = r["Date"], r["Heure"]   # Excel peut avoir converti les textes en date / heure
                    dt = dt.strftime("%d/%m/%Y") if hasattr(dt, "strftime") else str(dt)
                    hr = hr.strftime("%H:%M") if hasattr(hr, "strftime") else str(hr)[:5]
                    k = (ws.title, r.get("Symbole") or "", dt, hr, r["Sens"])
                    anciens[k] = {c: r.get(c) for c in PERSO}
    d = df[df["strat"] == "Strat5669"].sort_values("entree")
    wb = Workbook()
    wb.remove(wb.active)
    gris, jaune = PatternFill("solid", fgColor="D9D9D9"), PatternFill("solid", fgColor="FFE699")
    for nom, syms in FEUILLES.items():
        ws = wb.create_sheet(nom)
        cols = (["Symbole"] if nom == "BTC" else []) + AUTO + PERSO
        ws.append(cols)
        for i, c in enumerate(cols, 1):
            cell = ws.cell(1, i)
            cell.font = Font(bold=True)
            cell.fill = jaune if c in PERSO else gris
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[cell.column_letter].width = {"Remarque": 40, "Si non, pourquoi": 22, "Résultat": 15,
                                                              "Fantôme ?": 18}.get(c, 12)
        for r in d[d["sym"].isin(syms)].itertuples():
            fant = {"présent": "non", "ABSENT (fantôme)": "OUI"}.get(r.historique, r.historique.replace("présent, devenu", "non, devenu").replace("pas encore exporté", "à vérifier"))
            ligne = {"Symbole": r.sym, "Date": f"{r.entree:%d/%m/%Y}", "Heure": f"{r.entree:%H:%M}",
                     "Sens": "long" if r.sens == "L" else "short", "Étiquette": r.etiq,
                     "Résultat": (r.raison or "en cours").replace(" (P&L de l'historique)", ""),
                     "P&L alerte $": None if r.pnl is None or pd.isna(r.pnl) else round(r.pnl), "Fantôme ?": fant}
            k = (nom, r.sym if nom == "BTC" else "", ligne["Date"], ligne["Heure"], ligne["Sens"])
            ligne.update({c: v for c, v in anciens.get(k, {}).items() if v not in (None, "")})
            ws.append([ligne.get(c) for c in cols])
        ws.freeze_panes = "A2"
        n = max(ws.max_row, 200)
        lettre = {c: ws.cell(1, i).column_letter for i, c in enumerate(cols, 1)}
        for c, liste in (("Pris ?", "oui,non"), ("Si non, pourquoi", RAISONS)):
            dv = DataValidation(type="list", formula1=f'"{liste}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{lettre[c]}2:{lettre[c]}{n}")
        e, fa = lettre["Étiquette"], lettre["Fantôme ?"]
        ws.conditional_formatting.add(f"{e}2:{e}{n}", FormulaRule(formula=[f'{e}2="GODLIKE"'], fill=PatternFill("solid", fgColor="FFD966")))
        ws.conditional_formatting.add(f"{fa}2:{fa}{n}", FormulaRule(formula=[f'{fa}2="OUI"'], fill=PatternFill("solid", fgColor="F4B084")))
        p = lettre["P&L alerte $"]
        ws.conditional_formatting.add(f"{p}2:{p}{n}", FormulaRule(formula=[f"AND(ISNUMBER({p}2),{p}2<0)"], font=Font(color="C00000")))
    wb.save(chemin)
    print(f"Fiche de suivi : {chemin}")


def lignes_suivi(df, anciens, numbers=False):
    """Lignes de chaque feuille : colonnes auto depuis les alertes + colonnes perso relues dans l'ancien fichier."""
    d = df[df["strat"] == "Strat5669"].sort_values("entree")
    res = {}
    for nom, syms in FEUILLES.items():
        cols = (["Symbole"] if nom == "BTC" or numbers else []) + AUTO + PERSO
        lignes = []
        for r in d[d["sym"].isin(syms)].itertuples():
            fant = {"présent": "non", "ABSENT (fantôme)": "OUI"}.get(r.historique, r.historique.replace("présent, devenu", "non, devenu").replace("pas encore exporté", "à vérifier"))
            ligne = {"Symbole": r.sym, "Date": f"{r.entree:%d/%m/%Y}", "Heure": f"{r.entree:%H:%M}",
                     "Sens": "long" if r.sens == "L" else "short", "Étiquette": r.etiq,
                     "Résultat": (r.raison or "en cours").replace(" (P&L de l'historique)", ""),
                     "P&L alerte $": None if r.pnl is None or pd.isna(r.pnl) else int(round(r.pnl)), "Fantôme ?": fant}
            k = (nom, r.sym if nom == "BTC" or numbers else "", ligne["Date"], ligne["Heure"], ligne["Sens"])
            ligne.update({c: v for c, v in anciens.get(k, {}).items() if v not in (None, "")})
            lignes.append([ligne.get(c) for c in cols])
        res[nom] = (cols, lignes)
    return res


def en_nombre(v):
    """« 29223.25 », « -187,9 » tapés comme texte dans Numbers → nombre (pour que les totaux marchent)."""
    if isinstance(v, str):
        try:
            return float(v.strip().replace("\u202f", "").replace(" ", "").replace(",", ".").replace("$", ""))
        except ValueError:
            return v
    return v


def cle_ancienne(feuille, r):
    dt, hr = r.get("Date"), r.get("Heure")   # le tableur peut avoir converti les textes en date / heure
    dt = dt.strftime("%d/%m/%Y") if hasattr(dt, "strftime") else str(dt)
    hr = hr.strftime("%H:%M") if hasattr(hr, "strftime") else str(hr)[:5]
    return (feuille, r.get("Symbole") or "", dt, hr, r.get("Sens"))


def ecrire_suivi_numbers(df, chemin):
    """Même fiche, au format Numbers (bibliothèque numbers-parser) : menus locaux pour « Pris ? » et « Si non,
    pourquoi », en-têtes gris (auto) / jaunes (à remplir), colonnes perso conservées d'une fois sur l'autre."""
    from numbers_parser import Document, RGB
    anciens = {}
    if os.path.exists(chemin):
        for sh in Document(chemin).sheets:
            rows = sh.tables[0].rows(values_only=True)
            for row in rows[1:]:
                r = dict(zip(rows[0], row))
                if r.get("Date"):
                    a = {c: (None if r.get(c) == "—" else en_nombre(r.get(c))) for c in PERSO}
                    f = r.get("Fantôme ?")   # Jeremy y écrit parfois son propre constat : on le garde
                    if f not in (None, "", "non", "OUI", "à vérifier") and not str(f).startswith("non, devenu"):
                        a["Fantôme ?"] = f
                    anciens[cle_ancienne(sh.name, r)] = a
    feuilles = lignes_suivi(df, anciens, numbers=True)
    # numbers-parser donne à tous les tableaux la taille du dernier créé : on les crée donc tous de la même taille
    # (mêmes colonnes partout, Symbole compris ; autant de lignes que la feuille la plus longue).
    R = max(len(l) for _, l in feuilles.values()) + 1
    C = max(len(c) for c, _ in feuilles.values())
    doc = None
    for nom, (cols, lignes) in feuilles.items():
        if doc is None:
            doc = Document(num_header_cols=0, num_rows=R, num_cols=C)
            sh = doc.sheets[0]
            sh.name = nom
        else:
            doc.add_sheet(nom, "Trades", num_rows=R, num_cols=C)
            sh = doc.sheets[nom]
        t = sh.tables[0]
        t.name = "Trades"
        t.num_header_cols = 0
        if "auto" not in doc.styles:
            doc.add_style(name="auto", bg_color=RGB(217, 217, 217), bold=True, font_size=12.0)
            doc.add_style(name="perso", bg_color=RGB(255, 230, 153), bold=True, font_size=12.0)
            doc.add_style(name="godlike", bg_color=RGB(255, 217, 102))
            doc.add_style(name="fantome", bg_color=RGB(244, 176, 132), bold=True)
        for j, c in enumerate(cols):
            t.write(0, j, c, style="perso" if c in PERSO else "auto")
            t.col_width(j, {"Remarque": 260, "Si non, pourquoi": 150, "Résultat": 105, "Fantôme ?": 140,
                            "Date": 85}.get(c, 80))
        for i, ligne in enumerate(lignes, 1):
            for j, (c, v) in enumerate(zip(cols, ligne)):
                if v is None or v == "":
                    continue
                st = "godlike" if (c == "Étiquette" and v == "GODLIKE") else "fantome" if (c == "Fantôme ?" and str(v).lower().startswith("oui")) else None
                t.write(i, j, v, style=st) if st else t.write(i, j, v)
            # menu local : une cellule vide ne peut pas en porter un, d'où la valeur « — » (= pas encore rempli)
            for c, vals in (("Pris ?", ["—", "oui", "non"]), ("Si non, pourquoi", ["—"] + RAISONS.split(","))):
                j = cols.index(c)
                v = t.cell(i, j).value
                if v in (None, "") or v in vals:
                    if v in (None, ""):
                        t.write(i, j, "—")
                    t.set_cell_formatting(i, j, "popup", popup_values=vals, allow_none=True)
    doc.save(chemin)
    print(f"Fiche de suivi : {chemin}")


# --- 3. affichage ------------------------------------------------------------------------------------------------
df = pd.DataFrame(trades)
if df.empty:
    sys.exit("Aucune alerte d'entrée Strat5669 trouvée.")
print("TRADES SIGNALÉS EN DIRECT (journal d'alertes)\n")
for tr in trades:
    pnl = "en cours" if tr["sortie"] is None else ("P&L ?" if tr["pnl"] is None else f"{tr['pnl']:+,.0f} $")
    print(f"{tr['strat'][:14]:14} {tr['entree']:%d/%m %H:%M} {tr['sym']:7}{tr['tf']:>4} {tr['sens']} {tr['etiq']:8} @ {tr['px']:<10} "
          f"→ {tr['raison'] or '—':<16} {pnl:>10}   historique : {tr['historique']}")


def stats(sous, titre):
    fini = sous[sous["pnl"].notna()]
    if fini.empty:
        return
    g, p = fini.loc[fini["pnl"] > 0, "pnl"].sum(), -fini.loc[fini["pnl"] <= 0, "pnl"].sum()
    print(f"{titre:30} {len(sous):3} trades · {len(fini):3} terminés · {100 * (fini['pnl'] > 0).mean():5.1f} % gagnants · "
          f"PF {g / p if p else float('inf'):5.2f} · net {fini['pnl'].sum():+,.0f} $")


print("\nBILAN EN DIRECT")
for (st, sym, tf), s in df.groupby(["strat", "sym", "tf"]):
    stats(s, f"{st[:14]} {sym} {tf} tous")
    for etiq, s2 in s.groupby("etiq"):
        stats(s2, f"{st[:14]} {sym} {tf} {etiq}")
print("\nFANTÔMES (trades en direct comparés à l'historique)")
print(df.groupby(["strat", "etiq", "historique"]).size().to_string())

# Sens inverse : trades de l'historique, pendant que l'alerte tournait, sans aucune alerte d'entrée en direct.
print("\nNÉS APRÈS COUP (dans l'historique, jamais signalés en direct)")
nes = 0
for (st, sym, tf), s in df.groupby(["strat", "sym", "tf"]):
    t_ids = logs[logs["ID de l'alerte"].isin(set(s["id"]))]["t"]
    deb, fin = t_ids.min(), logs[logs["ID de l'alerte"].isin(set(s["id"]))]["fin_journal"].max()
    dernier = t_ids.max()   # après le dernier message de l'alerte : né après coup… ou alerte arrêtée
    couverts = [t for mx, t in exports.get((st, sym, tf), []) if mx >= deb]
    if not couverts:
        continue
    t = max(couverts, key=len)
    e = t[t["Type"].str.startswith("Entrer") & (t["Date et heure"] >= deb) & (t["Date et heure"] <= fin)]
    for r in e.itertuples():
        proche = s[((s["entree"] - r[3]).abs() <= pd.Timedelta(minutes=15)) & (s["sens"] == r.Signal[0])]
        if proche.empty:
            nes += 1
            doute = f" — après le dernier message de l'alerte ({dernier:%d/%m %H:%M}) : vérifier qu'elle était active" if r[3] > dernier else ""
            print(f"{st[:14]:14} {r[3]:%d/%m %H:%M} {sym:7}{tf:>4} {r.Signal:16} historique : P&L {r[8]:+,.0f} $ "
                  f"(journal du {deb:%d/%m %H:%M} au {fin:%d/%m %H:%M}){doute}")
if not nes:
    print("aucun")
if verrou:
    # Listes à coller dans Strat5669verrou.pine (réglage « Trades vécus en direct »), une par stratégie / symbole.
    os.makedirs(verrou, exist_ok=True)
    for (st, sym, tf), s in df.groupby(["strat", "sym", "tf"]):
        ids = set(s["id"])
        t_ids = logs[logs["ID de l'alerte"].isin(ids)]["t"]
        lignes = [f"DEBUT {t_ids.min():%Y-%m-%d %H:%M}"]
        lignes += [f"{r.entree:%Y-%m-%d %H:%M} {r.sens} {'G' if r.etiq == 'GODLIKE' else '5'}" for r in s.sort_values("entree").itertuples()]
        lignes.append(f"FIN {t_ids.max():%Y-%m-%d %H:%M}")
        nom = os.path.join(verrou, f"verrou_{st}_{sym.replace('!', '')}_{tf}.txt")
        open(nom, "w", encoding="utf-8").write("\n".join(lignes) + "\n")
        print(f"Liste verrou : {nom} ({len(s)} trades)")
if suivi:
    (ecrire_suivi_numbers if suivi.endswith(".numbers") else ecrire_suivi)(df, suivi)
if sortie:
    df.to_csv(sortie, index=False)
    print(f"\nTable écrite dans {sortie}")
